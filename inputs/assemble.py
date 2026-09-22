# -*- coding: utf-8 -*-
"""Assemble the final dataset: merge agent enrichment CSVs, compute 2026/27 opening dates and statuses,
merge Milan routing, apply score calibration overrides, verify webcams, validate, write the workbook.
Run from work/:  PYTHONUTF8=1 python assemble.py [--no-webcam-check]
Inputs:  ../candidates.csv, enrich_*.csv, routes.csv, overrides.csv (optional: id,field,value,reason)
Outputs: ../alpine_ski_resorts_2026_27.xlsx, ../resorts.csv, assembled_debug.csv, excluded.csv, validation_report.txt"""
import truststore; truststore.inject_into_ssl()
import csv, glob, sys, re, datetime as dt, math, time
from collections import Counter, OrderedDict
import requests
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

CHECK_WEBCAM = "--no-webcam-check" not in sys.argv
COUNTRY_ORDER = {"FR": 0, "IT": 1, "AT": 2, "CH": 3, "DE": 4, "SI": 5}
YEAR_ROUND = {"Hintertux Glacier", "Zermatt"}
# rows whose 2026/27 "published" date came only from aggregators (bergfex/skiinfo/skiresort) or provisional press round-ups
WEAK_PUBLISHED_IDS = {"117", "118", "121", "193", "35", "36", "38", "40", "44", "47", "48", "69",
                      "199", "200", "201", "205", "206", "209"}  # CH2 rows with skiresort.com season entries only
OUT_COLS = ["resort_name", "local_name", "country", "region", "ski_area", "type", "opening_date_2026_27", "opening_date_status",
            "max_elevation_m", "piste_km", "ticket_price_min", "ticket_price_max", "currency", "price_status",
            "driving_distance_km_from_milan", "driving_time_h_from_milan", "road_head", "freeride_score_0_5", "ski_touring_score_0_5",
            "score_notes", "webcam_url", "comments", "sources", "confidence"]

def rd(path):
    return list(csv.DictReader(open(path, encoding="utf-8-sig")))

def num(x):
    if x is None: return None
    s = str(x).strip().replace("€", "").replace("CHF", "").replace(",", ".").replace("'", "")
    if s == "" or s.lower() in ("nan", "none", "n/a", "tbd", "-"): return None
    m = re.search(r"-?\d+(\.\d+)?", s)
    return float(m.group()) if m else None

def parse_date(s):
    s = (s or "").strip()
    if not s or s.upper() == "TBD": return None
    for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y", "%Y/%m/%d"):
        try: return dt.datetime.strptime(s[:10], fmt).date()
        except ValueError: pass
    return None

def shift_to_2026(d):
    """2025/26 actual date -> same weekday one year later (+364 days)."""
    return d + dt.timedelta(days=364)

def half(x):
    v = num(x)
    if v is None: return None
    v = round(v * 2) / 2
    return max(0.0, min(5.0, v))

cands = {r["id"]: r for r in rd("../candidates.csv")}
enr = {}
for p in sorted(glob.glob("enrich_*.csv")):
    for r in rd(p):
        if r.get("id"): enr[r["id"].strip()] = r
routes = {r["id"]: r for r in rd("routes.csv")} if glob.glob("routes.csv") else {}
overrides = rd("overrides.csv") if glob.glob("overrides.csv") else []

missing = [i for i in cands if i not in enr]
rows, excluded, report = [], [], []
report.append(f"candidates={len(cands)} enriched={len(enr)} missing_enrichment={len(missing)}: {[cands[i]['resort_name'] for i in missing]}")

for cid, c in cands.items():
    e = enr.get(cid)
    if not e:
        continue
    comments = (e.get("comments") or "").strip()
    if comments.upper().startswith("EXCLUDE"):
        excluded.append(dict(id=cid, resort_name=c["resort_name"], country=c["country"], reason=comments)); continue
    row = OrderedDict()
    row["id"] = cid
    row["resort_name"] = c["resort_name"]
    row["local_name"] = (e.get("local_name") or c.get("local_name") or "").strip()
    if row["local_name"].lower() == row["resort_name"].lower(): row["local_name"] = ""
    row["country"] = c["country"]
    row["region"] = c["region"]
    row["ski_area"] = (e.get("ski_area") or c["ski_area"]).strip()
    t = (e.get("type") or "").strip()
    row["type"] = {"glacier": "Glacier", "glacier + non-glacier": "Glacier + non-glacier", "non-glacier": "Non-glacier"}.get(t.lower(), t)
    # opening date
    pub = parse_date(e.get("opening_2026_27_published")); act = parse_date(e.get("opening_2025_26_actual"))
    weak_no_actual = False
    if cid in WEAK_PUBLISHED_IDS and pub and act:
        comments = (comments + f" Aggregator/press lists a provisional 2026/27 opening of {pub.isoformat()}, not confirmed by the operator; date shown is the 2025/26 opening shifted.").strip()
        pub = None
    elif cid in WEAK_PUBLISHED_IDS and pub and not act:
        comments = (comments + " 2026/27 date is a resort-supplied season entry on an aggregator (skiresort.com/bergfex); no operator publication found and no 2025/26 date available.").strip()
        weak_no_actual = True
    if row["resort_name"] in YEAR_ROUND:
        row["opening_date_2026_27"] = pub or dt.date(2026, 10, 1); row["opening_date_status"] = "year_round_glacier"
    elif pub:
        row["opening_date_2026_27"] = pub; row["opening_date_status"] = "confirmed_2026_27"
    elif act:
        row["opening_date_2026_27"] = shift_to_2026(act); row["opening_date_status"] = "estimated_from_2025_26"
    else:
        row["opening_date_2026_27"] = "TBD"; row["opening_date_status"] = "TBD"
    row["_opening_2025_26_actual"] = act.isoformat() if act else ""
    row["_opening_source_note"] = (e.get("opening_source_note") or "").strip()
    row["max_elevation_m"] = int(round(num(e.get("max_elevation_m")))) if num(e.get("max_elevation_m")) is not None else None
    row["piste_km"] = int(round(num(e.get("piste_km")))) if num(e.get("piste_km")) is not None else None
    row["_piste_km_basis"] = (e.get("piste_km_basis") or "").strip()
    pmin, pmax = num(e.get("ticket_price_min")), num(e.get("ticket_price_max"))
    if pmin is not None and pmax is None: pmax = pmin
    if pmax is not None and pmin is None: pmin = pmax
    if pmin is not None and pmax is not None and pmin > pmax: pmin, pmax = pmax, pmin
    row["ticket_price_min"], row["ticket_price_max"] = pmin, pmax
    row["currency"] = "CHF" if c["country"] == "CH" else "EUR"
    ps = (e.get("price_status") or "").strip()
    row["price_status"] = "published_2026_27" if ps.startswith("published") else ("estimated_from_2025_26" if pmin is not None else "")
    # driving
    rt = routes.get(cid, {})
    dk, hh = num(rt.get("distance_km")), num(rt.get("hours"))
    row["driving_distance_km_from_milan"] = int(round(dk)) if dk is not None else None
    row["driving_time_h_from_milan"] = round(hh, 2) if hh is not None else None
    row["_route_status"] = rt.get("status", ""); row["_route_via"] = rt.get("via", ""); row["_route_hits"] = rt.get("hits_primary", ""); row["_route_modes"] = rt.get("modes", "")
    row["road_head"] = (e.get("road_head") or "").strip()
    row["_base_lat"], row["_base_lon"] = (e.get("base_lat") or "").strip(), (e.get("base_lon") or "").strip()
    row["freeride_score_0_5"] = half(e.get("freeride_score_0_5"))
    row["ski_touring_score_0_5"] = half(e.get("ski_touring_score_0_5"))
    row["score_notes"] = (e.get("score_notes") or "").strip()
    row["webcam_url"] = (e.get("webcam_url") or "").strip()
    row["comments"] = comments
    row["sources"] = (e.get("sources") or "").strip()
    conf = (e.get("confidence") or "").strip().lower()
    row["confidence"] = conf if conf in ("high", "medium", "low") else "medium"
    if weak_no_actual: row["confidence"] = "low"
    rows.append(row)

# overrides (manual corrections & score calibration): id, field, value, reason
for o in overrides:
    for row in rows:
        if row["id"] == o["id"].strip():
            f, v = o["field"].strip(), o["value"].strip()
            if f in ("freeride_score_0_5", "ski_touring_score_0_5"): row[f] = half(v)
            elif f in ("max_elevation_m", "piste_km", "driving_distance_km_from_milan"): row[f] = int(round(num(v))) if num(v) is not None else None
            elif f in ("ticket_price_min", "ticket_price_max", "driving_time_h_from_milan"): row[f] = num(v)
            elif f == "opening_date_2026_27":
                d = parse_date(v); row[f] = d if d else "TBD"
            elif f == "comments_append": row["comments"] = (row["comments"] + " " + v).strip()
            elif f == "sources_append": row["sources"] = (row["sources"] + " | " + v).strip(" |")
            elif f == "score_notes_append": row["score_notes"] = (row["score_notes"] + " " + v).strip()
            else: row[f] = v

# route comments
for row in rows:
    st = row["_route_status"]
    if st.startswith("detour_forced") and row["_route_hits"]:
        row["comments"] = (row["comments"] + f" Milan route forced via {row['_route_via']} to avoid closed {row['_route_hits'].replace(';', ', ')}.").strip()
    elif st == "alternative_used" and row["_route_hits"]:
        row["comments"] = (row["comments"] + f" Milan route uses OSRM alternative avoiding closed {row['_route_hits'].replace(';', ', ')}.").strip()
    if row["_route_modes"]:
        row["comments"] = (row["comments"] + f" Route includes car train/ferry segment ({row['_route_modes']}).").strip()
    if st.startswith("UNRESOLVED"):
        row["comments"] = (row["comments"] + f" ROUTING UNRESOLVED: {st}.").strip()

# derive statuses / confidence consistency
for row in rows:
    if row["opening_date_status"] in ("estimated_from_2025_26", "TBD") or row["price_status"] == "estimated_from_2025_26":
        if row["confidence"] == "high": row["confidence"] = "medium"
    est = sum([row["opening_date_status"] in ("estimated_from_2025_26", "TBD"), row["price_status"] != "published_2026_27",
               row["max_elevation_m"] is None, row["piste_km"] is None, row["ticket_price_min"] is None])
    if est >= 3: row["confidence"] = "low"

# webcam verification
if CHECK_WEBCAM:
    HD = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/124 Safari/537.36"}
    cache = {}
    for row in rows:
        u = row["webcam_url"]
        if not u: continue
        if u not in cache:
            code = None
            for attempt in range(3):
                try:
                    r = requests.get(u, headers=HD, timeout=45, allow_redirects=True); code = r.status_code; break
                except Exception as ex:
                    code = f"ERR {type(ex).__name__}"; time.sleep(5)
            cache[u] = code
        if cache[u] != 200:
            row["comments"] = (row["comments"] + f" Webcam URL dropped: failed verification at build time ({cache[u]}).").strip()
            row["webcam_url"] = ""
    report.append(f"webcam check: {sum(1 for v in cache.values() if v == 200)}/{len(cache)} URLs returned 200")

# validation
dups = [k for k, v in Counter((r["resort_name"], r["country"]) for r in rows).items() if v > 1]
report.append(f"duplicates: {dups}")
for row in rows:
    probs = []
    d = row["opening_date_2026_27"]
    if d == "TBD" and row["opening_date_status"] != "TBD": probs.append("date TBD but status not TBD")
    if isinstance(d, dt.date) and row["opening_date_status"] == "TBD": probs.append("date set but status TBD")
    if row["max_elevation_m"] is None or not (1000 <= row["max_elevation_m"] <= 4000): probs.append(f"elevation {row['max_elevation_m']}")
    if row["piste_km"] is None or row["piste_km"] <= 0: probs.append(f"piste_km {row['piste_km']}")
    for f in ("freeride_score_0_5", "ski_touring_score_0_5"):
        if row[f] is None: probs.append(f"{f} missing")
    if row["ticket_price_min"] is None: probs.append("price missing")
    if row["driving_distance_km_from_milan"] is None: probs.append("driving missing")
    if row["type"] not in ("Glacier", "Glacier + non-glacier", "Non-glacier"): probs.append(f"type '{row['type']}'")
    if row["_route_status"].startswith("UNRESOLVED"): probs.append("routing unresolved")
    if probs: report.append(f"CHECK {row['id']} {row['resort_name']} ({row['country']}): " + "; ".join(probs))

# sort
def sort_key(r):
    d = r["opening_date_2026_27"]
    return (COUNTRY_ORDER.get(r["country"], 9), 0 if isinstance(d, dt.date) else 1, d if isinstance(d, dt.date) else dt.date(2100, 1, 1), r["resort_name"])
rows.sort(key=sort_key)

# qa
qa = []
qa.append(("Rows total", len(rows)))
for cc in ["FR", "IT", "AT", "CH", "DE", "SI"]: qa.append((f"Rows {cc}", sum(1 for r in rows if r["country"] == cc)))
for t in ["Glacier", "Glacier + non-glacier", "Non-glacier"]: qa.append((f"Type {t}", sum(1 for r in rows if r["type"] == t)))
for s in ["confirmed_2026_27", "estimated_from_2025_26", "year_round_glacier", "TBD"]:
    n = sum(1 for r in rows if r["opening_date_status"] == s); qa.append((f"Opening status {s}", n))
n_conf = sum(1 for r in rows if r["opening_date_status"] == "confirmed_2026_27"); n_est = sum(1 for r in rows if r["opening_date_status"] == "estimated_from_2025_26")
qa.append(("Share confirmed_2026_27 vs estimated_from_2025_26 (% confirmed of the two)", round(100 * n_conf / max(1, n_conf + n_est), 1)))
for s in ["published_2026_27", "estimated_from_2025_26"]: qa.append((f"Price status {s}", sum(1 for r in rows if r["price_status"] == s)))
qa.append(("Blank webcam_url", sum(1 for r in rows if not r["webcam_url"])))
for cf in ["high", "medium", "low"]: qa.append((f"Confidence {cf}", sum(1 for r in rows if r["confidence"] == cf)))
qa.append(("Rows routed with forced detour (closed pass avoided)", sum(1 for r in rows if r["_route_status"].startswith("detour"))))
qa.append(("Excluded after research", len(excluded)))

NOTES = [
 "Inclusion: lift-served resort meeting at least one of (1) >=60 km marked pistes on the standard local day pass, (2) lift-served glacier skiing, (3) FWT/FWQ venue or consistently cited top freeride/ski-touring destination, (4) top lift >=2,800 m. One row per resort (own base village, own lift base, marketed as a destination); parent linked area in ski_area. Candidate discovery from skiresort.com (all 1,158 Alpine areas scraped and filtered) plus Bergfex/FWT cross-checks. Slovenia: no resort met any criterion, so it has 0 rows; Germany qualifies only via the Oberstdorf-Kleinwalsertal two-country pass and the Zugspitze.",
 "Glacier classification: 'Glacier' = early-season and core identity is glacier skiing (Hintertux, Stubai, Kaunertal, Pitztal, Moelltal, Kitzsteinhorn, Stelvio, Presena/Tonale); 'Glacier + non-glacier' = lift-served glacier terrain exists but most of the area is non-glacier (Zermatt, Saas-Fee, Tignes, Val d'Isere, Les 2 Alpes, Soelden, Engelberg, Cervinia, Val Senales); 'Non-glacier' = no lift-served glacier terrain.",
 "Opening dates: today is 2026-09-22. A 2026/27 date published by the operator or tourist office -> confirmed_2026_27 (many resorts publish these with their 2026/27 tariffs; dated pre-opening weekends count as the first day, with the continuous opening in comments). Otherwise the actual 2025/26 first lift day shifted +364 days to the equivalent 2026 weekday -> estimated_from_2025_26; 2026/27 dates that only appear on aggregators or in provisional press round-ups are noted in comments but not treated as confirmed (where no 2025/26 date exists either, the aggregator date is used and confidence is set to low). Year-round glaciers (Hintertux, Zermatt) carry the published start of the 2026/27 winter product -> year_round_glacier. TBD only where no winter operation exists (Passo dello Stelvio, summer glacier).",
 "Prices: adult one-day pass in the peak window (Christmas-New Year, February), local resort pass rather than mega-pass unless only one pass exists (basis noted in comments); published 2026/27 tariffs -> published_2026_27, else 2025/26 peak price -> estimated_from_2025_26. Dynamic-pricing resorts show a peak range (lower bound = documented peak-Saturday price such as the Blick 24 Jan 2026 comparison or the operator's 'from' price, upper bound = list/window price) with 'dynamic pricing' in comments; five dynamic-pricing resorts publish no usable figure and are left blank. EUR for FR/IT/AT/DE, CHF for CH.",
 "Driving: OSRM public demo server (car profile, OpenStreetMap data), origin Piazza del Duomo Milan, free-flow travel time (OSRM has no traffic model; weekday 06:00 departure is therefore equivalent to free flow). Distances are to the main base parking, or to the road head for car-free resorts (Taesch for Zermatt, Lauterbrunnen for Wengen/Muerren, valley stations for Aletsch Arena). Winter closures were enforced by detecting the closed-pass summits on the route geometry and forcing detours: brief list (Stelvio, Gavia, Umbrail, Great St Bernard pass, Furka, Grimsel, Susten, Nufenen, Oberalp, Klausen, Spluegen, San Bernardino pass, Flueela, Albula, Galibier, Iseran, Petit St Bernard, Bonette, Cayolle, Agnel, Lombarde, Timmelsjoch, Silvretta, Grossglockner) plus standard winter closures (Gotthard pass, Lukmanier, Forcola di Livigno, Izoard, Mont Cenis, Colombiere, Joux Plane, Croix de Fer, Glandon, Madeleine, Roselend, Sarenne, Allos, Champs, Couillole, Noyer, Penser Joch, Fedaia, Manghen, Finestre, San Carlo, Pfitscher Joch, Staller Sattel, Pramollo from Pontebba, Soelk, Col de la Croix, Sanetsch, Bassachaux, Joux Verte, Erbe, Duran, Hahntennjoch, Furkajoch, Pragel, Glaubenberg, Lech-Warth road). Tunnels and open passes per the brief were left available; car-train segments (Loetschberg, Vereina, Furka) count as driving time when OSRM uses them and are flagged in comments. Forced detours are noted in comments.",
 "Scores: freeride and ski-touring 0-5 in half steps per the brief's rubric, proposed per resort from FWT/FWQ venue history, guidebook/off-piste literature and Camptocamp/Gulliver/SAC route density, then calibrated by the coordinator against the anchor resorts (5.0: Chamonix, Verbier, Zermatt, Val d'Isere, La Grave, Engelberg, St. Anton; touring 5.0: Chamonix, Zermatt, Verbier, Engelberg, Andermatt). Scores are judgements, not measurements.",
 "Confidence: high = date, elevation, km and price from official sources; medium = one or more fields estimated (all estimated_from_2025_26 rows are at most medium); low = several fields estimated, missing or conflicting.",
 "Known limitations: 2026/27 dates and prices are mostly not yet published in September 2026, so most rows are estimates from 2025/26; OSRM free-flow times run slower than Google Maps on mountain roads (typically +10-20%); webcam URLs were verified for HTTP 200 at build time only; piste-km figures follow the lift companies' own claims, which differ from independently measured lengths.",
]

# write workbook
wb = Workbook()
ws = wb.active; ws.title = "resorts"
ws.append(OUT_COLS)
for row in rows:
    out = []
    for cname in OUT_COLS:
        v = row[cname]
        if isinstance(v, dt.date): out.append(v)
        else: out.append(v if v is not None else "")
    ws.append(out)
for c in ws[1]: c.font = Font(bold=True)
ws.freeze_panes = "A2"
ws.auto_filter.ref = ws.dimensions
date_col = OUT_COLS.index("opening_date_2026_27") + 1
for r in range(2, ws.max_row + 1):
    cell = ws.cell(row=r, column=date_col)
    if isinstance(cell.value, dt.date): cell.number_format = "YYYY-MM-DD"
    ws.cell(row=r, column=OUT_COLS.index("driving_time_h_from_milan") + 1).number_format = "0.00"
    for cname in ("freeride_score_0_5", "ski_touring_score_0_5"):
        ws.cell(row=r, column=OUT_COLS.index(cname) + 1).number_format = "0.0"
    for cname in ("ticket_price_min", "ticket_price_max"):
        ws.cell(row=r, column=OUT_COLS.index(cname) + 1).number_format = "0.00"
widths = {"resort_name": 30, "local_name": 18, "country": 8, "region": 26, "ski_area": 34, "type": 20, "opening_date_2026_27": 14, "opening_date_status": 22,
          "max_elevation_m": 12, "piste_km": 9, "ticket_price_min": 10, "ticket_price_max": 10, "currency": 9, "price_status": 22,
          "driving_distance_km_from_milan": 14, "driving_time_h_from_milan": 12, "road_head": 20, "freeride_score_0_5": 10, "ski_touring_score_0_5": 10,
          "score_notes": 60, "webcam_url": 45, "comments": 70, "sources": 70, "confidence": 11}
for i, cname in enumerate(OUT_COLS, start=1): ws.column_dimensions[get_column_letter(i)].width = widths.get(cname, 14)

wn = wb.create_sheet("notes")
wn.append(["Notes"]); wn["A1"].font = Font(bold=True)
for n in NOTES:
    wn.append(["- " + n])
wn.column_dimensions["A"].width = 160
for r in range(2, wn.max_row + 1): wn.cell(row=r, column=1).alignment = Alignment(wrap_text=True, vertical="top")

wq = wb.create_sheet("qa")
wq.append(["Metric", "Value"]); wq["A1"].font = Font(bold=True); wq["B1"].font = Font(bold=True)
for k, v in qa: wq.append([k, v])
wq.append([]); wq.append(["Low-confidence rows", ""]); wq.cell(row=wq.max_row, column=1).font = Font(bold=True)
for r in rows:
    if r["confidence"] == "low": wq.append([f"{r['resort_name']} ({r['country']})", r["comments"][:200]])
wq.append([]); wq.append(["Excluded after research", ""]); wq.cell(row=wq.max_row, column=1).font = Font(bold=True)
for x in excluded: wq.append([f"{x['resort_name']} ({x['country']})", x["reason"][:300]])
wq.column_dimensions["A"].width = 70; wq.column_dimensions["B"].width = 100
wb.save("../alpine_ski_resorts_2026_27.xlsx")

def csv_val(cname, v):
    if isinstance(v, dt.date): return v.isoformat()
    if v is None: return ""
    if cname == "driving_time_h_from_milan": return f"{v:.2f}"
    if cname in ("freeride_score_0_5", "ski_touring_score_0_5"): return f"{v:.1f}"
    return v
with open("../resorts.csv", "w", newline="", encoding="utf-8") as f:
    w = csv.writer(f); w.writerow(OUT_COLS)
    for row in rows:
        w.writerow([csv_val(c, row[c]) for c in OUT_COLS])
dbg_cols = ["id"] + OUT_COLS + [k for k in rows[0].keys() if k.startswith("_")] if rows else []
with open("assembled_debug.csv", "w", newline="", encoding="utf-8") as f:
    w = csv.writer(f); w.writerow(dbg_cols)
    for row in rows: w.writerow([(v.isoformat() if isinstance(v, dt.date) else ("" if v is None else v)) for v in (row.get(c) for c in dbg_cols)])
with open("excluded.csv", "w", newline="", encoding="utf-8") as f:
    w = csv.DictWriter(f, fieldnames=["id", "resort_name", "country", "reason"]); w.writeheader(); w.writerows(excluded)
with open("validation_report.txt", "w", encoding="utf-8") as f:
    f.write("\n".join(report) + "\n\nQA\n" + "\n".join(f"{k}: {v}" for k, v in qa) + "\n")
print("\n".join(report)); print("\nQA"); print("\n".join(f"{k}: {v}" for k, v in qa))
